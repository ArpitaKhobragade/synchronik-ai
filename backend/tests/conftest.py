"""Test fixtures: isolated SQLite DB, deterministic fake embedder/LLM/reranker.

No network, no model downloads, no API key: the full pipeline (ingestion →
hybrid retrieval → gate → generation → citations → persistence) runs offline.
"""
from __future__ import annotations

import hashlib
import io
import os

import numpy as np
import pytest

os.environ.update({"ENVIRONMENT": "test", "LOG_JSON": "false", "GEMINI_API_KEY": "", "ADMIN_API_KEY": "test-admin",
                   "RATE_LIMIT_CHAT_PER_MIN": "1000", "RATE_LIMIT_UPLOAD_PER_MIN": "1000", "MIN_SEMANTIC_SCORE": "0.2",
                   "LOAD_CURATED_DOCS": "false"})  # tests build their own knowledge base

from backend.core.config import get_settings  # noqa: E402
from backend.services.embeddings import EmbeddingProvider, set_embedder  # noqa: E402
from backend.services.llm import LLMProvider, set_llm  # noqa: E402
from backend.services.reranker import Reranker, set_reranker  # noqa: E402
from backend.services.vector_store import tokenize  # noqa: E402


class HashEmbedder(EmbeddingProvider):
    """Bag-of-words hashed into 256 dims: similar wording → similar vectors."""

    name, dim = "test-hash", 256

    def embed(self, texts):
        out = np.zeros((len(texts), self.dim), dtype="float32")
        for i, t in enumerate(texts):
            for tok in tokenize(t):
                out[i, int(hashlib.md5(tok.encode()).hexdigest(), 16) % self.dim] += 1
        norms = np.linalg.norm(out, axis=1, keepdims=True)
        return out / np.clip(norms, 1e-9, None)


class OverlapReranker(Reranker):
    name = "test-overlap"

    def score(self, query, passages):
        q = set(tokenize(query))
        return [len(q & set(tokenize(p))) / max(1, len(q)) for p in passages]


class FakeLLM(LLMProvider):
    model = "fake-llm"
    vision_model = "fake-vision"

    def __init__(self):
        self.calls = []

    def _reply(self, turns):
        last = turns[-1].text
        if "<sources>" in last:
            return "Attendance is marked per lecture by the faculty [1]. Edits after the cut-off need HOD approval [1]."
        return "standalone rewritten query about attendance"

    def generate(self, system, turns, *, max_tokens=None):
        self.calls.append((system, turns))
        return self._reply(turns)

    def stream(self, system, turns):
        self.calls.append((system, turns))
        for word in self._reply(turns).split(" "):
            yield word + " "

    def describe_image(self, system, question, image_bytes, mime):
        return "The screenshot shows the Attendance Entry screen."


@pytest.fixture()
def app_env(tmp_path, monkeypatch):
    from backend.db import session as dbs
    from backend.services import vector_store
    from backend.core.security import limiter

    # TEST_DATABASE_URL (e.g. the PostgreSQL service in CI) runs the same suite against PostgreSQL;
    # every test starts from empty tables. Default: a throwaway SQLite file per test.
    pg_url = os.environ.get("TEST_DATABASE_URL")
    monkeypatch.setenv("DATABASE_URL", pg_url or f"sqlite:///{tmp_path / 'test.db'}")
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    get_settings.cache_clear()
    dbs.reset_engine()
    if pg_url:
        from backend.db.models import Base

        Base.metadata.drop_all(dbs.get_engine())
    set_embedder(HashEmbedder())
    set_reranker(OverlapReranker())
    llm = FakeLLM()
    set_llm(llm)
    vector_store.knowledge_index = vector_store.KnowledgeIndex()
    # modules that imported the singleton by name
    import backend.services.retrieval as r, backend.services.ingestion as ing, backend.api.routes.misc as misc
    import backend.services.analytics as an
    for mod in (r, ing, misc, an):
        monkeypatch.setattr(mod, "knowledge_index", vector_store.knowledge_index, raising=False)
    dbs.init_db()
    vector_store.knowledge_index.load_or_build()
    limiter.reset()
    yield {"llm": llm, "index": vector_store.knowledge_index}
    dbs.reset_engine()
    get_settings.cache_clear()


@pytest.fixture()
def client(app_env):
    from fastapi.testclient import TestClient
    from backend.main import create_app

    with TestClient(create_app()) as c:
        yield c


ATTENDANCE_PAGE = """
<html><head><title>Attendance · Help</title></head><body><nav>Home Menu</nav><main>
<h2>Attendance entry</h2><p>Attendance is marked per lecture by the faculty from the Attendance Entry screen.
Select the course, date and period, then mark each student present or absent and save.</p>
<h2>Editing past attendance</h2><p>Edits after the cut-off date require HOD approval. The HOD unlocks the date
from Attendance Settings and the faculty can then correct the entry.</p>
</main><footer>Copyright</footer></body></html>"""

LIBRARY_PAGE = """
<html><head><title>Library · Help</title></head><body><main>
<h2>Library circulation</h2><p>Books are issued and returned at the circulation desk. Late returns accrue a fine
per day configured in Library Settings.</p></main></body></html>"""


def make_docx(paragraphs: list[tuple[str, str | None]]) -> bytes:
    from docx import Document

    d = Document()
    for text, style in paragraphs:
        d.add_paragraph(text, style=style) if style else d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(pages: list[str]) -> bytes:
    """Minimal valid PDF with one text line per page (no extra dependencies)."""
    objs, kids = [], []
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(None)  # pages, filled later
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        content_id = len(objs)
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
                    f"/Resources << /Font << /F1 3 0 R >> >> >>".encode())
        kids.append(len(objs))
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>".encode()
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


@pytest.fixture()
def seeded(app_env):
    from backend.services.document_processing import extract_html
    from backend.services.ingestion import upsert_web_page

    for url, html in (("http://help.synchronik.co.in/attendance.html", ATTENDANCE_PAGE),
                      ("http://help.synchronik.co.in/library.html", LIBRARY_PAGE)):
        title, sections = extract_html(html)
        upsert_web_page(url, title, sections, "help_portal")
    return app_env
